"""Bounded CSV/XLSX parsing and atomic, idempotent imports through record service."""
import csv
import hashlib
import io
import json
import math
import re
import zipfile
from datetime import date, datetime
from decimal import Decimal,InvalidOperation
from openpyxl.utils.exceptions import InvalidFileException
from xml.etree.ElementTree import ParseError

from fastapi import HTTPException
from sqlalchemy import select
from app.api.deps import current_tenant_id, current_user_id
from app.models.relational import Form, SystemSetting
from app.services import records
from app.services.form_runtime import _latest_form_version

MAX_BYTES = 5 * 1024 * 1024
MAX_ROWS = 500
MAX_COLUMNS = 100


def parse_file(filename, content):
    if len(content) > MAX_BYTES:
        raise HTTPException(422, '文件不能超过 5 MB')
    try:
        if filename.lower().endswith('.csv'):
            text = content.decode('utf-8-sig')
            try: dialect = csv.Sniffer().sniff(text[:8192], delimiters=',;\t')
            except csv.Error: dialect = csv.excel
            rows=[]
            for row in csv.reader(io.StringIO(text),dialect=dialect,strict=True):
                if not any(value not in (None,'') for value in row):continue
                rows.append(row)
                if len(rows)>MAX_ROWS+1 or len(row)>MAX_COLUMNS:raise HTTPException(422,'最多导入 500 行、100 列')
        elif filename.lower().endswith('.xlsx'):
            from openpyxl import load_workbook
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                if len(archive.infolist()) > 1000 or sum(item.file_size for item in archive.infolist()) > 25 * 1024 * 1024:
                    raise HTTPException(422, 'Excel 解压后超过上限')
            book = load_workbook(io.BytesIO(content), read_only=True, data_only=False, keep_links=False)
            try:
                rows = []
                for row in book.worksheets[0].iter_rows():
                    if len(rows) > MAX_ROWS or len(row) > MAX_COLUMNS:
                        raise HTTPException(422, '最多导入 500 行、100 列，Excel 使用第一个工作表')
                    if any(cell.data_type == 'f' for cell in row):
                        raise HTTPException(422, 'Excel 包含公式，请先粘贴为值再导入')
                    rows.append([cell.value.isoformat() if isinstance(cell.value, (date, datetime)) else cell.value for cell in row])
            finally: book.close()
        else:
            raise HTTPException(422, '仅支持 UTF-8 CSV 和 .xlsx 文件')
    except (UnicodeError, csv.Error, ValueError, zipfile.BadZipFile, KeyError, IndexError, OSError, InvalidFileException, ParseError):
        raise HTTPException(422, '文件格式无效；CSV 请使用 UTF-8 编码，Excel 请使用 .xlsx')
    rows = [row for row in rows if any(value not in (None, '') for value in row)]
    if len(rows) < 2 or len(rows) > MAX_ROWS + 1 or len(rows[0]) > MAX_COLUMNS:
        raise HTTPException(422, '请提供标题行及 1～500 行数据，最多 100 列')
    headers = [str(value or '').strip() for value in rows[0]]
    if any(not value for value in headers) or len(set(headers)) != len(headers):
        raise HTTPException(422, '列标题不能为空或重复')
    result = []
    for number, row in enumerate(rows[1:], 2):
        if len(row) > len(headers) and any(value not in (None, '') for value in row[len(headers):]):
            raise HTTPException(422, f'第 {number} 行的列数超过标题行')
        result.append([*row[:len(headers)], *[None] * max(0, len(headers)-len(row))])
    return headers, result


async def context(db, user, form_id):
    form = await db.get(Form, form_id)
    if not form or form.tenant_id != current_tenant_id(user):
        raise HTTPException(404, 'Form not found')
    await records.require_form_permission(db, user, form_id, 'import')
    await records.require_form_permission(db, user, form_id, 'create')
    version = await _latest_form_version(db, form.tenant_id, form.id)
    if not version:
        raise HTTPException(409, '请先发布表单配置，再批量导入')
    fields = await records.runtime_fields(db, form)
    from app.core.permissions import allowed_form_fields
    from app.services.form_rules import states
    allowed = await allowed_form_fields(user, form_id, 'create', fields, db)
    state = states(fields, {})
    derived = {item.get('targetField') for field in fields for item in (field.ui_config or {}).get('assignmentRules') or []}
    for field in fields:
        relation = (field.ui_config or {}).get('relation_config') or {}
        if isinstance(relation, dict): derived.update((relation.get('backfill') or {}).keys())
    writable = [field for field in fields if not field.archived and field.field_name in allowed and field.field_name not in derived
                and not state[field.field_name]['readonly'] and not state[field.field_name].get('generated')]
    return form, version, fields, writable


def convert(value, field):
    if value in (None, ''): return None
    kind = field.field_type.lower()
    try:
        if kind in {'int', 'integer'}:
            number = Decimal(str(value))
            if isinstance(value,bool) or not number.is_finite() or number!=number.to_integral_value():raise ValueError()
            result=int(number)
            if not -(2**31)<=result<2**31:raise ValueError()
            return result
        if kind in {'number', 'float', 'decimal'}:
            number = float(value)
            if isinstance(value, bool) or not math.isfinite(number): raise ValueError()
            return number
        if kind == 'boolean':
            lowered = str(value).strip().lower()
            if lowered in {'true', '1', '是'}: return True
            if lowered in {'false', '0', '否'}: return False
            raise ValueError()
        if kind == 'json': return json.loads(value) if isinstance(value, str) else value
        if kind in {'date', 'datetime'}:
            raw = str(value)
            if kind=='date':
                try:return date.fromisoformat(raw).isoformat()
                except ValueError:return datetime.fromisoformat(raw.replace('Z','+00:00')).date().isoformat()
            return datetime.fromisoformat(raw.replace('Z','+00:00')).isoformat()
        return str(value)
    except (ValueError, TypeError, OverflowError,InvalidOperation):
        raise HTTPException(422, f'字段「{field.label}」不能转换为 {field.field_type}')


async def preview(db, user, form_id, filename, content):
    _, version, _, writable = await context(db, user, form_id)
    headers, rows = parse_file(filename, content)
    mapping = {};used=set()
    for index, name in enumerate(headers):
        exact = [field for field in writable if field.field_name == name]
        matches = exact or [field for field in writable if field.label == name]
        if len(matches)==1 and matches[0].field_name not in used:
            mapping[str(index)]=matches[0].field_name;used.add(matches[0].field_name)
    return {'headers': headers, 'rows': rows[:10], 'row_count': len(rows), 'mapping': mapping, 'schema_version': version.version,
            'fields': [{'name': field.field_name, 'label': field.label, 'type': field.field_type} for field in writable]}


async def execute(db,user,form_id,filename,content,mapping,request_key,schema_version):
    from app.services.form_option_sources import read_snapshot
    async with read_snapshot():return await _execute(db,user,form_id,filename,content,mapping,request_key,schema_version)


async def _execute(db, user, form_id, filename, content, mapping, request_key, schema_version):
    if not re.fullmatch(r'[a-zA-Z0-9-]{16,64}', request_key):
        raise HTTPException(422, '导入请求标识无效，请重新预览')
    form, version, fields, writable = await context(db, user, form_id)
    headers, rows = parse_file(filename, content)
    from app.services.field_uniqueness import lock
    await lock(db,form.tenant_id)
    await db.refresh(form)
    form,version,fields,writable=await context(db,user,form_id)
    if not isinstance(mapping, dict): raise HTTPException(422, '列映射无效')
    by_name = {field.field_name: field for field in writable}
    try:
        mapped = [(int(index), name) for index, name in mapping.items() if name]
        if not mapped or any(index < 0 or index >= len(headers) or name not in by_name for index, name in mapped) or len({name for _, name in mapped}) != len(mapped):
            raise ValueError()
    except (ValueError, TypeError):
        raise HTTPException(422, '请选择有效且不重复的列映射；只读、自动编码和无权限字段不可导入')
    digest = hashlib.sha256(content + json.dumps(mapping, sort_keys=True).encode()).hexdigest()
    key = f'form_import:{form.tenant_id}:{current_user_id(user)}:{form_id}:{request_key}'
    receipt = await db.scalar(select(SystemSetting).where(SystemSetting.key == key))
    if receipt:
        if receipt.value.get('digest') != digest: raise HTTPException(409, '此请求标识已用于其他文件，请重新预览')
        return {**receipt.value['result'], 'replayed': True}
    if schema_version != version.version:
        raise HTTPException(409, '表单配置已发布新版本，请重新预览并检查列映射')
    from app.services.form_rules import evaluate
    from app.api.form_engine.validation import _validate_record_data
    from app.services.form_schemas import DynamicRecordCreate
    from app.services.form_application_lifecycle import enabled
    managed = await enabled(db, form)
    ids = []
    for row_number, row in enumerate(rows, 2):
        try:
            payload = {name: convert(row[index], by_name[name]) for index, name in mapped}
            for field in writable:
                if field.field_name not in payload and field.default_value not in (None, ''):
                    payload[field.field_name] = convert(field.default_value, field)
            checked = await evaluate(db, user, fields, payload, enforce=True)
            _validate_record_data(checked['fields'], checked['data'])
            result = await records.create_record(form_id, DynamicRecordCreate(data=payload), db, user, defer_commit=True, draft=managed, source='bulk-import')
            ids.append(result['data']['id'])
        except HTTPException as exc:
            await db.rollback()
            raise HTTPException(exc.status_code, {'message': f'第 {row_number} 行导入失败，整批未写入', 'row': row_number, 'detail': exc.detail})
    result = {'count': len(ids), 'record_ids': ids, 'schema_version': version.version, 'draft': managed, 'replayed': False}
    db.add(SystemSetting(key=key, value={'digest': digest, 'result': result}, description='Atomic form import receipt', updated_by=str(current_user_id(user))))
    await db.commit()
    return result
